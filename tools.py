from pathlib import Path
import os
import shutil
import subprocess
import tempfile

import fitz  # PyMuPDF
from PIL import Image
from pypdf import PdfReader, PdfWriter
import pdfplumber
from openpyxl import Workbook
from pptx import Presentation
from pptx.util import Inches

# =========================================================
# BASIC CONFIGURATION
# =========================================================

BASE_DIR = Path(__file__).resolve().parent
OUTPUT_DIR = BASE_DIR / "outputs"

def ensure_output_folder():
    OUTPUT_DIR.mkdir(parents=True, exist_ok=True)
    return OUTPUT_DIR

def safe_output_path(filename):
    ensure_output_folder()
    return OUTPUT_DIR / Path(filename).name

# =========================================================
# LIBREOFFICE
# =========================================================

def find_libreoffice():
    possible_paths = [
        r"C:\Program Files\LibreOffice\program\soffice.exe",
        r"C:\Program Files (x86)\LibreOffice\program\soffice.exe",
        "/usr/bin/libreoffice",
        "/usr/bin/soffice",
        "/Applications/LibreOffice.app/Contents/MacOS/soffice",
    ]
    for path in possible_paths:
        if os.path.exists(path):
            return path
    return shutil.which("soffice") or shutil.which("libreoffice")

def run_libreoffice(input_file, output_dir):
    soffice = find_libreoffice()
    if not soffice:
        raise RuntimeError("LibreOffice is not installed or could not be found.")

    input_file, output_dir = Path(input_file), Path(output_dir)
    output_dir.mkdir(parents=True, exist_ok=True)

    command = [
        soffice, "--headless", "--convert-to", "pdf",
        "--outdir", str(output_dir), str(input_file)
    ]
    result = subprocess.run(command, capture_output=True, text=True, timeout=120)
    if result.returncode != 0:
        raise RuntimeError(f"LibreOffice conversion failed:\n{result.stderr}")

    output_file = output_dir / f"{input_file.stem}.pdf"
    if not output_file.exists():
        raise RuntimeError("LibreOffice finished, but the PDF output was not created.")
    return output_file

# =========================================================
# 1. PDF -> IMAGE
# =========================================================

def pdf_to_image(input_file, output_dir=None, image_format="png"):
    input_file = Path(input_file)
    output_dir = Path(output_dir or OUTPUT_DIR)
    output_dir.mkdir(parents=True, exist_ok=True)

    image_format = image_format.lower()
    if image_format not in ("png", "jpg", "jpeg"):
        image_format = "png"

    pdf = fitz.open(input_file)
    generated_files = []
    try:
        for page_number, page in enumerate(pdf, start=1):
            pixmap = page.get_pixmap(matrix=fitz.Matrix(2, 2), alpha=False)
            ext = "jpg" if image_format in ("jpg", "jpeg") else "png"
            output_file = output_dir / f"{input_file.stem}_page_{page_number}.{ext}"
            if ext == "jpg":
                pixmap.save(str(output_file), output="jpeg")
            else:
                pixmap.save(str(output_file))
            generated_files.append(output_file)
    finally:
        pdf.close()
    return generated_files

# =========================================================
# 2. IMAGE -> PDF
# =========================================================

def image_to_pdf(input_files, output_file=None):
    if isinstance(input_files, (str, Path)):
        input_files = [input_files]
    input_files = [Path(f) for f in input_files]
    if not input_files:
        raise ValueError("No image files were provided.")

    output_file = Path(output_file or safe_output_path("images.pdf"))
    images = []
    try:
        for file in input_files:
            if not file.exists():
                raise FileNotFoundError(f"Image not found: {file}")
            img = Image.open(file)
            images.append(img.convert("RGB") if img.mode not in ("RGB", "L") else img.copy())

        images[0].save(
            output_file, "PDF", resolution=100.0,
            save_all=True, append_images=images[1:]
        )
    finally:
        for img in images:
            img.close()
    return output_file

# =========================================================
# 3. MERGE PDF
# =========================================================

def merge_pdfs(input_files, output_file=None):
    if isinstance(input_files, (str, Path)):
        input_files = [input_files]
    input_files = [Path(f) for f in input_files]
    if len(input_files) < 2:
        raise ValueError("At least two PDF files are required.")

    output_file = Path(output_file or safe_output_path("merged.pdf"))
    writer = PdfWriter()
    for file in input_files:
        if not file.exists():
            raise FileNotFoundError(f"PDF not found: {file}")
        reader = PdfReader(str(file))
        for page in reader.pages:
            writer.add_page(page)

    with open(output_file, "wb") as output:
        writer.write(output)
    return Path(output_file)

# =========================================================
# 4. SPLIT PDF
# =========================================================

def split_pdf(input_file, output_dir=None):
    input_file = Path(input_file)
    output_dir = Path(output_dir or OUTPUT_DIR)
    output_dir.mkdir(parents=True, exist_ok=True)

    reader = PdfReader(str(input_file))
    generated_files = []
    for page_number, page in enumerate(reader.pages, start=1):
        output_file = output_dir / f"{input_file.stem}_page_{page_number}.pdf"
        writer = PdfWriter()
        writer.add_page(page)
        with open(output_file, "wb") as output:
            writer.write(output)
        generated_files.append(output_file)
    return generated_files

# =========================================================
# 5. COMPRESS PDF
# =========================================================

def compress_pdf(input_file, output_file=None):
    input_file = Path(input_file)
    output_file = Path(output_file or safe_output_path(f"{input_file.stem}_compressed.pdf"))
    document = fitz.open(input_file)
    try:
        document.save(
            output_file, garbage=4, clean=True, deflate=True,
            deflate_images=True, deflate_fonts=True
        )
    finally:
        document.close()
    return output_file

# =========================================================
# 6. PDF -> TEXT
# =========================================================

def pdf_to_text(input_file, output_file=None):
    input_file = Path(input_file)
    output_file = Path(output_file or safe_output_path(f"{input_file.stem}.txt"))
    pdf = fitz.open(input_file)
    try:
        with open(output_file, "w", encoding="utf-8") as text_file:
            for page_number, page in enumerate(pdf, start=1):
                text_file.write(f"\n===== PAGE {page_number} =====\n\n")
                text_file.write(page.get_text())
    finally:
        pdf.close()
    return output_file

# =========================================================
# 7. LOCK PDF
# =========================================================

def lock_pdf(input_file, password, output_file=None):
    if not password:
        raise ValueError("Password cannot be empty.")
    input_file = Path(input_file)
    output_file = Path(output_file or safe_output_path(f"{input_file.stem}_protected.pdf"))

    reader = PdfReader(str(input_file))
    writer = PdfWriter()
    for page in reader.pages:
        writer.add_page(page)
    writer.encrypt(user_password=password, owner_password=password)

    with open(output_file, "wb") as output:
        writer.write(output)
    return output_file

# =========================================================
# 8, 9, 10. OFFICE CONVERSIONS -> PDF
# =========================================================

def word_to_pdf(input_file, output_dir=None):
    return run_libreoffice(Path(input_file), output_dir or OUTPUT_DIR)

def ppt_to_pdf(input_file, output_dir=None):
    return run_libreoffice(Path(input_file), output_dir or OUTPUT_DIR)

def excel_to_pdf(input_file, output_dir=None):
    return run_libreoffice(Path(input_file), output_dir or OUTPUT_DIR)

# =========================================================
# 11. PDF -> WORD
# =========================================================

def pdf_to_word(input_file, output_file=None):
    from pdf2docx import Converter
    input_file = Path(input_file)
    output_file = Path(output_file or safe_output_path(f"{input_file.stem}.docx"))
    converter = Converter(str(input_file))
    try:
        converter.convert(str(output_file))
    finally:
        converter.close()
    return output_file

# =========================================================
# 12. PDF -> POWERPOINT
# =========================================================

def pdf_to_ppt(input_file, output_file=None):
    input_file = Path(input_file)
    output_file = Path(output_file or safe_output_path(f"{input_file.stem}.pptx"))

    pdf = fitz.open(input_file)
    presentation = Presentation()
    presentation.slide_width = Inches(13.333333)
    presentation.slide_height = Inches(7.5)
    blank_layout = presentation.slide_layouts[6]
    temporary_files = []

    try:
        for page_number, page in enumerate(pdf, start=1):
            temp_file = Path(tempfile.gettempdir()) / f"docutools_pdf_page_{os.getpid()}_{page_number}.png"
            pixmap = page.get_pixmap(matrix=fitz.Matrix(2, 2), alpha=False)
            pixmap.save(str(temp_file))
            temporary_files.append(temp_file)

            slide = presentation.slides.add_slide(blank_layout)
            slide.shapes.add_picture(
                str(temp_file), 0, 0,
                width=presentation.slide_width, height=presentation.slide_height
            )
        presentation.save(output_file)
    finally:
        pdf.close()
        for file in temporary_files:
            try:
                if file.exists():
                    file.unlink()
            except OSError:
                pass
    return output_file

# =========================================================
# 13. PDF -> EXCEL
# =========================================================

def pdf_to_excel(input_file, output_file=None):
    input_file = Path(input_file)
    output_file = Path(output_file or safe_output_path(f"{input_file.stem}.xlsx"))

    workbook = Workbook()
    workbook.remove(workbook.active)
    found_table = False

    with pdfplumber.open(input_file) as pdf:
        for page_number, page in enumerate(pdf.pages, start=1):
            tables = page.extract_tables()
            for table_number, table in enumerate(tables, start=1):
                if not table:
                    continue
                found_table = True
                worksheet = workbook.create_sheet(title=f"Page{page_number}_Table{table_number}"[:31])
                for row in table:
                    if row is not None:
                        worksheet.append(["" if cell is None else str(cell).strip() for cell in row])

    if not found_table:
        worksheet = workbook.create_sheet(title="Extracted Data")
        worksheet["A1"] = "No structured tables were detected in this PDF."

    workbook.save(output_file)
    return output_file

# =========================================================
# 14. CROP PDF
# =========================================================

def crop_pdf(input_file, output_file=None, left=0, top=0, right=0, bottom=0):
    """Crop PDF pages using margins in PDF points (72 points = 1 inch)."""
    input_file = Path(input_file)
    output_file = Path(output_file or safe_output_path(f"{input_file.stem}_cropped.pdf"))

    try:
        left, top, right, bottom = float(left), float(top), float(right), float(bottom)
    except (TypeError, ValueError):
        raise ValueError("Crop values must be numbers.")

    if min(left, top, right, bottom) < 0:
        raise ValueError("Crop values cannot be negative.")

    pdf = fitz.open(input_file)
    try:
        for page in pdf:
            rect = page.rect
            nl, nt = rect.x0 + left, rect.y0 + top
            nr, nb = rect.x1 - right, rect.y1 - bottom
            if nl >= nr or nt >= nb:
                raise ValueError("Crop margins are too large for the page size.")
            page.set_cropbox(fitz.Rect(nl, nt, nr, nb))
        pdf.save(output_file, garbage=4, deflate=True)
    finally:
        pdf.close()
    return output_file

# =========================================================
# 15. SIGN PDF
# =========================================================

def sign_pdf(input_file, signature_file, output_file=None, page_number=1, x=50, y=50, width=150, height=60):
    """Add a signature image to a PDF page at (x, y) coordinates."""
    input_file, signature_file = Path(input_file), Path(signature_file)
    if not signature_file.exists():
        raise FileNotFoundError("Signature image was not found.")

    output_file = Path(output_file or safe_output_path(f"{input_file.stem}_signed.pdf"))

    try:
        page_number, x, y, width, height = int(page_number), float(x), float(y), float(width), float(height)
    except (TypeError, ValueError):
        raise ValueError("Signature position and size must be numeric.")

    if page_number < 1:
        raise ValueError("Page number must be at least 1.")
    if width <= 0 or height <= 0:
        raise ValueError("Signature size must be greater than zero.")

    pdf = fitz.open(input_file)
    try:
        if page_number > len(pdf):
            raise ValueError(f"PDF has only {len(pdf)} page(s).")
        page = pdf[page_number - 1]
        rect = fitz.Rect(x, y, x + width, y + height) & page.rect
        if rect.is_empty:
            raise ValueError("Signature position is outside the page.")
        page.insert_image(rect, filename=str(signature_file), keep_proportion=True, overlay=True)
        pdf.save(output_file, garbage=4, deflate=True)
    finally:
        pdf.close()
    return output_file

# =========================================================
# 16. ADD PAGE NUMBERS
# =========================================================

def add_page_numbers(input_file, output_file=None, position="bottom-center", start_number=1, font_size=10):
    """Add page numbers to every PDF page."""
    input_file = Path(input_file)
    output_file = Path(output_file or safe_output_path(f"{input_file.stem}_numbered.pdf"))

    valid_positions = {"bottom-left", "bottom-center", "bottom-right", "top-left", "top-center", "top-right"}
    if position not in valid_positions:
        raise ValueError("Invalid page number position.")

    try:
        start_number, font_size = int(start_number), float(font_size)
    except (TypeError, ValueError):
        raise ValueError("Page number settings are invalid.")
    if font_size <= 0:
        raise ValueError("Font size must be greater than zero.")

    pdf = fitz.open(input_file)
    try:
        margin = 25
        for index, page in enumerate(pdf, start=0):
            text = str(start_number + index)
            rect = page.rect

            if position == "bottom-left":
                box, align = fitz.Rect(margin, rect.height - 35, 100, rect.height - 10), 0
            elif position == "bottom-center":
                box, align = fitz.Rect(0, rect.height - 35, rect.width, rect.height - 10), 1
            elif position == "bottom-right":
                box, align = fitz.Rect(rect.width - 100, rect.height - 35, rect.width - margin, rect.height - 10), 2
            elif position == "top-left":
                box, align = fitz.Rect(margin, 10, 100, 35), 0
            elif position == "top-center":
                box, align = fitz.Rect(0, 10, rect.width, 35), 1
            else:  # top-right
                box, align = fitz.Rect(rect.width - 100, 10, rect.width - margin, 35), 2

            page.insert_textbox(
                box, text, fontsize=font_size, fontname="helv",
                color=(0, 0, 0), align=align, overlay=True
            )
        pdf.save(output_file, garbage=4, deflate=True)
    finally:
        pdf.close()
    return output_file

# =========================================================
# TOOL LIST & SELF TEST
# =========================================================

TOOL_NAMES = {
    "pdf-to-image": "PDF → Image",
    "image-to-pdf": "Image → PDF",
    "merge-pdf": "Merge PDF",
    "split-pdf": "Split PDF",
    "compress-pdf": "Compress PDF",
    "pdf-to-text": "PDF → Text",
    "lock-pdf": "Lock PDF",
    "word-to-pdf": "Word → PDF",
    "powerpoint-to-pdf": "PowerPoint → PDF",
    "excel-to-pdf": "Excel → PDF",
    "pdf-to-word": "PDF → Word",
    "pdf-to-powerpoint": "PDF → PowerPoint",
    "pdf-to-excel": "PDF → Excel",
    "crop-pdf": "Crop PDF",
    "sign-pdf": "Sign PDF",
    "add-page-numbers": "Add Page Numbers",
}

if __name__ == "__main__":
    print("=" * 55)
    print("DocuTools Backend")
    print("=" * 55)
    print("\nAvailable tools:\n")
    for key, name in TOOL_NAMES.items():
        print(f"✓ {key:25} -> {name}")
    print("\nTotal tools:", len(TOOL_NAMES))
    print("\nOutput folder:", OUTPUT_DIR)
    libreoffice = find_libreoffice()
    print("\nLibreOffice:", f"✓ Found: {libreoffice}" if libreoffice else "✗ LibreOffice not found")
    print("\nBackend loaded successfully.")